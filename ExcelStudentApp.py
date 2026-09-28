import os
from datetime import date, datetime

import customtkinter as ctk
from tkinter import filedialog, messagebox, ttk
from openpyxl import load_workbook


ctk.set_appearance_mode("light")
ctk.set_default_color_theme("blue")


class ExcelStudentApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("Excel Student Reader")
        self.geometry("1500x750")
        self.minsize(1100, 600)

        self.excel_file_path = None
        self.worksheet_name = None
        self.excel_row_map = {}

        # These names must match the headers in Excel row 1.
        self.columns = (
            "ID",
            "First Name",
            "Last Name",
            "Father",
            "Mother",
            "Address",
            "DOB",
            "Age",
            "Gender",
            "Session",
            "Class",
            "Roll",
        )

        self.create_style()
        self.create_top_bar()
        self.create_table()
        self.create_status_bar()

    def create_style(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except Exception:
            pass

        style.configure(
            "Treeview",
            rowheight=32,
            font=("Segoe UI", 10),
            background="white",
            fieldbackground="white",
            foreground="#111827",
        )
        style.configure(
            "Treeview.Heading",
            font=("Segoe UI", 10, "bold"),
            background="#E5E7EB",
            foreground="#111827",
            padding=8,
        )
        style.map(
            "Treeview",
            background=[("selected", "#2563EB")],
            foreground=[("selected", "white")],
        )

    def create_top_bar(self):
        top_frame = ctk.CTkFrame(self)
        top_frame.pack(fill="x", padx=10, pady=(10, 5))

        ctk.CTkButton(
            top_frame,
            text="Browse Excel",
            width=125,
            command=self.browse_excel,
        ).pack(side="left", padx=5, pady=10)

        ctk.CTkButton(
            top_frame,
            text="Import",
            width=95,
            command=self.import_excel,
        ).pack(side="left", padx=5, pady=10)

        ctk.CTkButton(
            top_frame,
            text="Edit Selected Row",
            width=150,
            command=self.open_edit_window,
        ).pack(side="left", padx=5, pady=10)

        ctk.CTkButton(
            top_frame,
            text="Refresh From Excel",
            width=150,
            command=self.import_excel,
        ).pack(side="left", padx=5, pady=10)

        ctk.CTkButton(
            top_frame,
            text="Clear Table",
            width=110,
            fg_color="#D97706",
            hover_color="#B45309",
            command=self.clear_table,
        ).pack(side="left", padx=5, pady=10)

        self.file_label = ctk.CTkLabel(
            top_frame,
            text="No Excel file selected",
            anchor="w",
        )
        self.file_label.pack(side="left", fill="x", expand=True, padx=15)

    def create_table(self):
        table_frame = ctk.CTkFrame(self)
        table_frame.pack(fill="both", expand=True, padx=10, pady=5)

        tree_container = ctk.CTkFrame(table_frame, fg_color="transparent")
        tree_container.pack(fill="both", expand=True, padx=5, pady=5)

        self.tree = ttk.Treeview(
            tree_container,
            columns=self.columns,
            show="headings",
            selectmode="browse",
        )

        column_widths = {
            "ID": 70,
            "First Name": 140,
            "Last Name": 140,
            "Father": 140,
            "Mother": 140,
            "Address": 220,
            "DOB": 110,
            "Age": 70,
            "Gender": 90,
            "Session": 100,
            "Class": 100,
            "Roll": 90,
        }

        for column in self.columns:
            self.tree.heading(column, text=column)
            self.tree.column(
                column,
                width=column_widths.get(column, 120),
                minwidth=60,
                anchor="center",
            )

        vertical_scrollbar = ttk.Scrollbar(
            tree_container,
            orient="vertical",
            command=self.tree.yview,
        )
        horizontal_scrollbar = ttk.Scrollbar(
            tree_container,
            orient="horizontal",
            command=self.tree.xview,
        )

        self.tree.configure(
            yscrollcommand=vertical_scrollbar.set,
            xscrollcommand=horizontal_scrollbar.set,
        )

        self.tree.grid(row=0, column=0, sticky="nsew")
        vertical_scrollbar.grid(row=0, column=1, sticky="ns")
        horizontal_scrollbar.grid(row=1, column=0, sticky="ew")

        tree_container.grid_rowconfigure(0, weight=1)
        tree_container.grid_columnconfigure(0, weight=1)

        self.tree.bind("<Double-1>", self.on_tree_double_click)

    def create_status_bar(self):
        self.status_label = ctk.CTkLabel(
            self,
            text="Ready",
            anchor="w",
            height=30,
        )
        self.status_label.pack(fill="x", padx=15, pady=(0, 5))

    def update_status(self, message):
        self.status_label.configure(text=message)

    def browse_excel(self):
        selected_file = filedialog.askopenfilename(
            title="Select Excel File",
            filetypes=[
                ("Excel Workbook", "*.xlsx"),
                ("Excel Macro-Enabled Workbook", "*.xlsm"),
            ],
        )

        if not selected_file:
            return

        self.excel_file_path = selected_file
        self.worksheet_name = None
        self.file_label.configure(text=os.path.basename(selected_file))
        self.update_status(f"Selected file: {selected_file}")

    def is_macro_workbook(self):
        return bool(
            self.excel_file_path
            and self.excel_file_path.lower().endswith(".xlsm")
        )

    def open_workbook(self):
        return load_workbook(
            self.excel_file_path,
            data_only=False,
            keep_vba=self.is_macro_workbook(),
        )

    @staticmethod
    def normalize_header(header):
        if header is None:
            return ""
        return str(header).strip()

    def get_worksheet_headers(self, worksheet):
        return [self.normalize_header(cell.value) for cell in worksheet[1]]

    def get_column_indexes(self, worksheet):
        excel_headers = self.get_worksheet_headers(worksheet)
        missing_columns = [
            column for column in self.columns if column not in excel_headers
        ]

        if missing_columns:
            return None, missing_columns

        column_indexes = {
            column: excel_headers.index(column) + 1 for column in self.columns
        }
        return column_indexes, []

    @staticmethod
    def display_value(value):
        if value is None:
            return ""
        if isinstance(value, (datetime, date)):
            return value.strftime("%d/%m/%Y")
        return value

    def import_excel(self):
        if not self.excel_file_path:
            messagebox.showwarning(
                "No File Selected",
                "Please browse and select an Excel file first.",
            )
            return

        workbook = None
        try:
            workbook = self.open_workbook()
            worksheet = workbook.active
            self.worksheet_name = worksheet.title

            column_indexes, missing_columns = self.get_column_indexes(worksheet)
            if missing_columns:
                messagebox.showerror(
                    "Missing Excel Columns",
                    "The Excel file is missing these columns:\n\n"
                    + ", ".join(missing_columns)
                    + "\n\nPlease check Excel row 1.",
                )
                return

            self.clear_table(show_message=False)
            imported_count = 0

            for excel_row_number in range(2, worksheet.max_row + 1):
                values = []
                for column in self.columns:
                    value = worksheet.cell(
                        row=excel_row_number,
                        column=column_indexes[column],
                    ).value
                    values.append(self.display_value(value))

                if not any(str(value).strip() for value in values):
                    continue

                tree_item = self.tree.insert("", "end", values=values)
                self.excel_row_map[tree_item] = excel_row_number
                imported_count += 1

            self.update_status(
                f"Imported {imported_count} rows from worksheet "
                f"'{self.worksheet_name}'."
            )
            messagebox.showinfo(
                "Import Completed",
                f"{imported_count} rows imported successfully.",
            )

        except PermissionError:
            messagebox.showerror(
                "File Access Error",
                "The Excel file cannot be accessed.\n\n"
                "Close the file in Microsoft Excel and try again.",
            )
        except Exception as error:
            messagebox.showerror(
                "Import Error",
                f"Could not import the Excel file.\n\n{error}",
            )
        finally:
            if workbook is not None:
                workbook.close()

    def on_tree_double_click(self, event):
        row_id = self.tree.identify_row(event.y)
        if not row_id:
            return
        self.tree.selection_set(row_id)
        self.tree.focus(row_id)
        self.open_edit_window()

    def open_edit_window(self):
        selected_items = self.tree.selection()
        if not selected_items:
            messagebox.showwarning(
                "No Row Selected",
                "Please select a row to edit.",
            )
            return

        selected_item = selected_items[0]
        current_values = self.tree.item(selected_item, "values")

        edit_window = ctk.CTkToplevel(self)
        edit_window.title("Edit Student")
        edit_window.geometry("700x680")
        edit_window.resizable(False, False)
        edit_window.transient(self)
        edit_window.grab_set()

        ctk.CTkLabel(
            edit_window,
            text="Edit Student Information",
            font=ctk.CTkFont(size=20, weight="bold"),
        ).pack(pady=(20, 10))

        form_frame = ctk.CTkScrollableFrame(
            edit_window,
            width=620,
            height=500,
        )
        form_frame.pack(fill="both", expand=True, padx=20, pady=10)

        edit_entries = {}

        for row_index, column in enumerate(self.columns):
            ctk.CTkLabel(
                form_frame,
                text=f"{column}:",
                width=130,
                anchor="w",
            ).grid(row=row_index, column=0, padx=10, pady=7, sticky="w")

            current_value = ""
            if row_index < len(current_values):
                current_value = str(current_values[row_index])

            if column == "Gender":
                gender_variable = ctk.StringVar(value=current_value)
                gender_frame = ctk.CTkFrame(form_frame, fg_color="transparent")
                gender_frame.grid(
                    row=row_index,
                    column=1,
                    padx=10,
                    pady=7,
                    sticky="w",
                )
                ctk.CTkRadioButton(
                    gender_frame,
                    text="Male",
                    variable=gender_variable,
                    value="Male",
                ).pack(side="left", padx=(0, 20))
                ctk.CTkRadioButton(
                    gender_frame,
                    text="Female",
                    variable=gender_variable,
                    value="Female",
                ).pack(side="left")
                edit_entries[column] = gender_variable
            else:
                entry = ctk.CTkEntry(form_frame, width=400)
                entry.grid(
                    row=row_index,
                    column=1,
                    padx=10,
                    pady=7,
                    sticky="ew",
                )
                entry.insert(0, current_value)
                edit_entries[column] = entry

        form_frame.grid_columnconfigure(1, weight=1)

        button_frame = ctk.CTkFrame(edit_window, fg_color="transparent")
        button_frame.pack(fill="x", padx=20, pady=(5, 20))

        ctk.CTkButton(
            button_frame,
            text="Update UI Only",
            width=180,
            fg_color="#D97706",
            hover_color="#B45309",
            command=lambda: self.update_ui_only(
                selected_item,
                edit_entries,
                edit_window,
            ),
        ).pack(side="left", padx=5, expand=True)

        ctk.CTkButton(
            button_frame,
            text="Save to Excel",
            width=180,
            fg_color="#15803D",
            hover_color="#166534",
            command=lambda: self.save_row_to_excel(
                selected_item,
                edit_entries,
                edit_window,
            ),
        ).pack(side="left", padx=5, expand=True)

        ctk.CTkButton(
            button_frame,
            text="Cancel",
            width=120,
            fg_color="#6B7280",
            hover_color="#4B5563",
            command=edit_window.destroy,
        ).pack(side="left", padx=5, expand=True)

    def get_edit_values(self, edit_entries):
        return [edit_entries[column].get().strip() for column in self.columns]

    @staticmethod
    def validate_values(values, parent_window):
        if not values[0]:
            messagebox.showwarning(
                "Missing ID",
                "Student ID cannot be empty.",
                parent=parent_window,
            )
            return False
        return True

    def update_ui_only(self, selected_item, edit_entries, edit_window):
        updated_values = self.get_edit_values(edit_entries)
        if not self.validate_values(updated_values, edit_window):
            return

        self.tree.item(selected_item, values=updated_values)
        self.update_status(
            "The selected row was updated in the UI only. "
            "The Excel file was not changed."
        )
        messagebox.showinfo(
            "UI Updated",
            "The selected row was updated inside the application only.\n\n"
            "The Excel workbook was not changed.",
            parent=edit_window,
        )
        edit_window.destroy()

    def save_row_to_excel(self, selected_item, edit_entries, edit_window):
        if not self.excel_file_path:
            messagebox.showwarning(
                "No Excel File",
                "There is no Excel file connected to this row.",
                parent=edit_window,
            )
            return

        if selected_item not in self.excel_row_map:
            messagebox.showwarning(
                "Excel Row Not Found",
                "The selected UI row is not linked to an Excel row.",
                parent=edit_window,
            )
            return

        updated_values = self.get_edit_values(edit_entries)
        if not self.validate_values(updated_values, edit_window):
            return

        workbook = None
        try:
            workbook = self.open_workbook()

            if self.worksheet_name and self.worksheet_name in workbook.sheetnames:
                worksheet = workbook[self.worksheet_name]
            else:
                worksheet = workbook.active

            column_indexes, missing_columns = self.get_column_indexes(worksheet)
            if missing_columns:
                messagebox.showerror(
                    "Missing Excel Columns",
                    "The Excel file is missing these columns:\n\n"
                    + ", ".join(missing_columns),
                    parent=edit_window,
                )
                return

            excel_row_number = self.excel_row_map[selected_item]

            for column, value in zip(self.columns, updated_values):
                worksheet.cell(
                    row=excel_row_number,
                    column=column_indexes[column],
                ).value = value

            workbook.save(self.excel_file_path)

            self.tree.item(selected_item, values=updated_values)
            self.update_status(
                f"Excel row {excel_row_number} was updated successfully."
            )
            messagebox.showinfo(
                "Saved to Excel",
                "The selected row was updated in the UI and saved "
                "to the Excel workbook.",
                parent=edit_window,
            )
            edit_window.destroy()

        except PermissionError:
            messagebox.showerror(
                "File Access Error",
                "The Excel file may currently be open.\n\n"
                "Close the Excel file and try again.",
                parent=edit_window,
            )
        except Exception as error:
            messagebox.showerror(
                "Save Error",
                f"Could not save the changes to Excel.\n\n{error}",
                parent=edit_window,
            )
        finally:
            if workbook is not None:
                workbook.close()

    def clear_table(self, show_message=True):
        for item in self.tree.get_children():
            self.tree.delete(item)

        self.excel_row_map.clear()
        self.update_status("The table is empty.")

        if show_message:
            messagebox.showinfo(
                "Table Cleared",
                "All rows were removed from the UI.\n\n"
                "The Excel workbook was not changed.",
            )


if __name__ == "__main__":
    app = ExcelStudentApp()
    app.mainloop()
